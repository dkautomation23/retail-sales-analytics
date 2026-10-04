"""The workbook builder on a small made-up year. No database."""
from datetime import date

from openpyxl import load_workbook

from analytics.export_excel import build


def rows():
    months = [date(2010, m, 1) for m in range(1, 13)] + [date(2011, m, 1) for m in range(1, 4)]
    return [(m, 1000.0 + i * 10, 100.0 + i, 10 + i) for i, m in enumerate(months)]


def test_derived_columns_are_formulas_and_raw_columns_are_values(tmp_path):
    path = tmp_path / "x.xlsx"
    build(rows(), 12345.0, date(2011, 3, 1)).save(path)
    ws = load_workbook(path)["Monthly"]
    assert ws["B2"].value == 1000.0 and ws["C2"].value == 100.0
    assert ws["E2"].value == "=B2-C2"
    assert ws["G14"].value.startswith("=IFERROR(E14/INDEX(")  # year on year looks 12 months back
    assert ws["I15"].value.startswith("=IF(A15=DATE(2011,3,1)")


def test_summary_has_a_check_cell_against_the_warehouse(tmp_path):
    path = tmp_path / "x.xlsx"
    build(rows(), 12345.0, date(2011, 3, 1)).save(path)
    sm = load_workbook(path)["Summary"]
    assert sm["B2"].value == 12345.0
    assert sm["B3"].value == "=ROUND(B1-B2,2)"
    assert sm["B5"].value.startswith("=INDEX(Monthly!A2:A16,MATCH(MAX(")
    assert sm["A9"].value == 2010 and sm["A10"].value == 2011
    assert sm["B9"].value.startswith("=SUMIFS(")


def test_no_formula_uses_single_quoted_text(tmp_path):
    """'' is not an empty string in Excel; a formula with it is rejected on open."""
    path = tmp_path / "x.xlsx"
    build(rows(), 12345.0, date(2011, 3, 1)).save(path)
    wb = load_workbook(path)
    for ws in wb:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    assert "''" not in cell.value, f"{ws.title}!{cell.coordinate}"
