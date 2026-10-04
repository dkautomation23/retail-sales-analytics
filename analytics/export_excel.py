"""A small Excel workbook for people who do not open a database: monthly revenue with live formulas.

    python -m analytics.export_excel      # writes bi/retail-sales-summary.xlsx

Monthly gross revenue, cancelled value and order count are values read from the warehouse.
Everything derived from them (net revenue, month on month and year on year change, the yearly
summary, the best month, the check against the warehouse total) is an Excel formula, so
changing a month recalculates the rest. Summary!B3 compares the formulas' total with the
warehouse's own net revenue and must read 0.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from .db import connect

OUT = Path(__file__).resolve().parent.parent / "bi" / "retail-sales-summary.xlsx"
MONTHLY = """
SELECT d.month_start, coalesce(s.gross, 0), coalesce(r.cancelled, 0), coalesce(s.orders, 0)
FROM (SELECT DISTINCT month_start FROM dim_date) d
LEFT JOIN (SELECT dd.month_start, sum(revenue) AS gross, count(DISTINCT invoice) AS orders
           FROM fact_sales f JOIN dim_date dd USING (date_key) GROUP BY 1) s USING (month_start)
LEFT JOIN (SELECT dd.month_start, sum(amount) AS cancelled
           FROM fact_returns f JOIN dim_date dd USING (date_key) GROUP BY 1) r USING (month_start)
ORDER BY d.month_start
"""
WAREHOUSE_NET = "SELECT (SELECT sum(revenue) FROM fact_sales) - (SELECT sum(amount) FROM fact_returns)"
HEADER = ["Month", "Gross revenue", "Cancelled", "Orders", "Net revenue", "Net MoM %", "Net YoY %", "Year", "Partial month"]
MONEY, PERCENT = '#,##0;[Red]-#,##0', '0.0%'


def build(rows: list, warehouse_net: float, partial) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Monthly"
    ws.append(HEADER)
    last = len(rows) + 1
    for i, (month, gross, cancelled, orders) in enumerate(rows, start=2):
        mom = f'=IFERROR(E{i}/E{i - 1}-1,"")' if i > 2 else ""
        yoy = f'=IFERROR(E{i}/INDEX($E$2:$E${last},MATCH(EDATE(A{i},-12),$A$2:$A${last},0))-1,"")'
        ws.append([month, float(gross), float(cancelled), int(orders), f"=B{i}-C{i}", mom, yoy, f"=YEAR(A{i})",
                   f'=IF(A{i}=DATE({partial.year},{partial.month},1),"partial","")'])
        ws[f"A{i}"].number_format = "mmm yyyy"
        for col in "BCE":
            ws[f"{col}{i}"].number_format = MONEY
        for col in "FG":
            ws[f"{col}{i}"].number_format = PERCENT
    table = Table(displayName="MonthlyTable", ref=f"A1:I{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "B2"
    ws.conditional_formatting.add(f"F2:G{last}", CellIsRule(operator="lessThan", formula=["0"],
                                                          fill=PatternFill("solid", bgColor="FFC7CE")))
    sm = wb.create_sheet("Summary")
    sm["A1"], sm["B1"] = "Net revenue, formulas total", f"=SUM(Monthly!E2:E{last})"
    sm["A2"], sm["B2"] = "Net revenue, warehouse (SQL)", warehouse_net
    sm["A3"], sm["B3"] = "Difference (must be 0)", "=ROUND(B1-B2,2)"
    sm["A5"], sm["B5"] = "Best month by net revenue", f"=INDEX(Monthly!A2:A{last},MATCH(MAX(Monthly!E2:E{last}),Monthly!E2:E{last},0))"
    sm["A6"], sm["B6"] = "Net revenue in that month", f"=MAX(Monthly!E2:E{last})"
    for col, text in zip("ABCDE", ("Year", "Gross revenue", "Cancelled", "Net revenue", "Cancellation rate")):
        sm[f"{col}8"] = text
        sm[f"{col}8"].font = Font(bold=True)
    years = sorted({r[0].year for r in rows})
    for n, year in enumerate(years, start=9):
        sm[f"A{n}"] = year
        sm[f"B{n}"] = f"=SUMIFS(Monthly!$B$2:$B${last},Monthly!$H$2:$H${last},A{n})"
        sm[f"C{n}"] = f"=SUMIFS(Monthly!$C$2:$C${last},Monthly!$H$2:$H${last},A{n})"
        sm[f"D{n}"] = f"=B{n}-C{n}"
        sm[f"E{n}"] = f'=IFERROR(C{n}/B{n},"")'
        sm[f"E{n}"].number_format = '0.00%'
        for col in "BCD":
            sm[f"{col}{n}"].number_format = MONEY
    for cell in ("B1", "B2", "B6"):
        sm[cell].number_format = MONEY
    sm["B5"].number_format = "mmm yyyy"
    sm.column_dimensions["A"].width = 30
    for col in "BCDE":
        sm.column_dimensions[col].width = 18
    return wb


def main() -> int:
    with connect() as conn:
        rows = conn.execute(MONTHLY).fetchall()
        net = float(conn.execute(WAREHOUSE_NET).fetchone()[0])
        partial = conn.execute("SELECT max(month_start) FROM dim_date").fetchone()[0]
    OUT.parent.mkdir(parents=True, exist_ok=True)
    build(rows, net, partial).save(OUT)
    print(f"wrote {OUT.name}: {len(rows)} months, warehouse net revenue {net:,.2f}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
